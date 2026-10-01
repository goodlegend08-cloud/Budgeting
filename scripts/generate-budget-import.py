#!/usr/bin/env python3
"""Build public/budget-import-2026.json from the budget spreadsheet.

Usage:
  python3 scripts/generate-budget-import.py sheet.xlsx [previous.json] [out.json]

If previous.json is given it is diffed against the new data, so rows whose
amount/date changed and budget limits that disappeared are emitted as
remove* entries — a re-import then updates an already-populated tracker
instead of duplicating rows.
"""

import json
import re
import sys
from datetime import date, datetime

from openpyxl import load_workbook

MONTH_NAMES = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12,
}

CATEGORIES = {
    "debts": ("Debts", "💸", "#f43f5e"),
    "savings": ("Savings", "🏦", "#0d9488"),
    "necessities": ("Necessities", "🧺", "#f59e0b"),
    "transport": ("Transport", "🚇", "#0ea5e9"),
    "shopping": ("Shopping", "🛍️", "#ec4899"),
    "fees": ("Fees & Documents", "📄", "#64748b"),
}


def map_item(name: str):
    key = name.strip().lower()
    if "utang" in key:
        return CATEGORIES["debts"]
    if key == "savings":
        return CATEGORIES["savings"]
    if key == "necessities":
        return CATEGORIES["necessities"]
    if key == "gasoline":
        return CATEGORIES["transport"]
    if "vac" in key or "washing" in key:
        return CATEGORIES["shopping"]
    if "notary" in key or "landbank" in key or "passport" in key:
        return CATEGORIES["fees"]
    return None


def find_label_column(rows) -> int:
    """The sheet's label column drifts (September uses C, other tabs use B)."""
    for row in rows[:40]:
        for index, value in enumerate(row[:6]):
            text = str(value).strip()
            if text.startswith("Block ") or text == "Item / Name":
                return index
    return 1


def parse_date(raw, month_num: int, year: int = 2026) -> str:
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    if raw:
        text = str(raw).strip()
        named = re.search(r"([A-Za-z]+)\s+(\d{1,2})", text)
        if named and named.group(1).lower() in MONTH_NAMES:
            month = MONTH_NAMES[named.group(1).lower()]
            return f"{year}-{month:02d}-{int(named.group(2)):02d}"
        iso = re.match(r"(\d{4})-(\d{2})-(\d{2})", text)
        if iso:
            return iso.group(0)
    return f"{year}-{month_num:02d}-15"


def next_month(month_num: int) -> tuple[int, int]:
    return (2027, 1) if month_num == 12 else (2026, month_num + 1)


def cell(row, index):
    return row[index] if index < len(row) else ""


def build(sheet_path: str) -> dict:
    wb = load_workbook(sheet_path, data_only=True)
    categories, transactions, budgets, salary = {}, [], {}, []

    for sheet_name in ("September", "October", "November", "December"):
        if sheet_name not in wb.sheetnames:
            continue
        month_num = MONTH_NAMES[sheet_name.lower()[:3]]
        rows = [
            [("" if value is None else value) for value in row[:8]]
            for row in wb[sheet_name].iter_rows(values_only=True)
        ]
        label_col = find_label_column(rows)
        date_col = label_col + 1
        amount_col = label_col + 2

        gross = deductions = None
        gross_parts = []
        for row in rows:
            label = str(cell(row, label_col)).strip()
            amount = cell(row, amount_col)
            if label == "Total Gross" and isinstance(amount, (int, float)):
                gross = float(amount)
            elif label == "Standard Deductions" and isinstance(amount, (int, float)):
                deductions = abs(float(amount))
            elif label in ("SG2 Basic", "PERA") and isinstance(amount, (int, float)):
                gross_parts.append(f"{label} ₱{float(amount):,.2f}")

        if gross:
            lines = (
                [{"name": "Standard deductions", "amount": round(deductions, 2)}]
                if deductions
                else []
            )
            salary.append(
                {
                    "date": f"2026-{month_num:02d}-15",
                    "label": f"{sheet_name} 2026 Salary",
                    "grossPay": round(gross, 2),
                    "allowances": [],
                    "deductions": lines,
                    "note": " + ".join(gross_parts),
                }
            )

        # Block 1 is salary; later blocks are spending. Block 4+ is "gastos
        # before the next salary", so it is planned for the following month.
        section = None
        for row in rows:
            label = str(cell(row, label_col)).strip()
            block = re.match(r"Block (\d+):", label)
            if block:
                section = int(block.group(1))
                continue
            if label.startswith("Bottom Line"):
                section = None
                continue
            if section is None or section == 1:
                continue
            if label in ("", "Item / Name") or label.startswith("Total "):
                continue

            amount = cell(row, amount_col)
            if not isinstance(amount, (int, float)) or float(amount) <= 0:
                continue
            mapped = map_item(label)
            if not mapped:
                print(f"!! unmapped: {sheet_name} / {label}")
                continue

            name, icon, color = mapped
            categories[name] = {
                "name": name,
                "type": "expense",
                "icon": icon,
                "color": color,
            }
            value = round(float(amount), 2)
            transactions.append(
                {
                    "date": parse_date(cell(row, date_col), month_num),
                    "amount": value,
                    "type": "expense",
                    "category": name,
                    "note": label,
                }
            )

            if section >= 4:
                year, following = next_month(month_num)
                budget_month = f"{year}-{following:02d}"
            else:
                budget_month = f"2026-{month_num:02d}"
            key = f"{budget_month}-{name}"
            budgets[key] = {
                "month": budget_month,
                "category": name,
                "limit": round(budgets.get(key, {}).get("limit", 0) + value, 2),
            }

    return {
        "version": 1,
        "source": "Google Sheet — Annual Budget Dashboard - 2026",
        "categories": list(categories.values()),
        "transactions": sorted(transactions, key=lambda t: (t["date"], t["note"])),
        "budgets": sorted(budgets.values(), key=lambda b: (b["month"], b["category"])),
        "salary": salary,
    }


def diff_removals(payload: dict, previous: dict) -> None:
    old_transactions = {
        (t["category"].lower(), t["note"].strip().lower()): t
        for t in previous.get("transactions", [])
    }
    changed = []
    for tx in payload["transactions"]:
        key = (tx["category"].lower(), tx["note"].strip().lower())
        old = old_transactions.get(key)
        if old and (old["amount"] != tx["amount"] or old["date"] != tx["date"]):
            changed.append({"category": tx["category"], "note": tx["note"]})
    payload["removeTransactions"] = changed

    keep = {(b["month"], b["category"].lower()) for b in payload["budgets"]}
    payload["removeBudgets"] = [
        {"month": b["month"], "category": b["category"]}
        for b in previous.get("budgets", [])
        if (b["month"], b["category"].lower()) not in keep
    ]


def main() -> None:
    sheet = sys.argv[1]
    previous_path = sys.argv[2] if len(sys.argv) > 2 else None
    out_path = sys.argv[3] if len(sys.argv) > 3 else "public/budget-import-2026.json"

    payload = build(sheet)
    if previous_path:
        try:
            with open(previous_path) as handle:
                diff_removals(payload, json.load(handle))
        except FileNotFoundError:
            pass

    with open(out_path, "w") as handle:
        json.dump(payload, handle, indent=2, ensure_ascii=False)
        handle.write("\n")

    spent = sum(t["amount"] for t in payload["transactions"])
    limited = sum(b["limit"] for b in payload["budgets"])
    print(f"categories:     {len(payload['categories'])}")
    print(f"transactions:   {len(payload['transactions'])} (₱{spent:,.2f})")
    print(f"budgets:        {len(payload['budgets'])} (₱{limited:,.2f})")
    print(f"salary records: {len(payload['salary'])}")
    print(
        "replace rows:   "
        f"{len(payload.get('removeTransactions', []))} transactions, "
        f"{len(payload.get('removeBudgets', []))} limits"
    )


if __name__ == "__main__":
    main()
